#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = 'sources/Paid_Media_Audit_Master_Findings_v18.xlsx'
OUT = 'FINAL_v_current/Paid_Media_Audit_Master_Findings_v20.xlsx'

wb = openpyxl.load_workbook(SRC)

# ---- style helpers matching existing palette ----
WHITE = 'FFFFFFFF'
def font(b=False, sz=11, color='FF000000'):
    return Font(bold=b, size=sz, color=color)
def fill(rgb):
    return PatternFill('solid', fgColor=rgb)
BANNER_FILL = fill('FF2E75B6'); BANNER_FONT = font(True, 14, WHITE)
SECT_FILL   = fill('FF8EA9DB'); SECT_FONT   = font(True, 11)
HDR_FILL    = fill('FF305496'); HDR_FONT    = font(True, 11, WHITE)
ALL_FILL    = fill('FFE2EFDA'); ALL_FONT    = font(True, 11)
NOTE_FONT   = font(False, 11, 'FF595959')
thin = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFTW = Alignment(horizontal='left', vertical='top', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')

def banner(ws, row, text, lastcol='J'):
    ws.merge_cells(f'A{row}:{lastcol}{row}')
    c = ws[f'A{row}']; c.value = text; c.font = BANNER_FONT; c.fill = BANNER_FILL; c.alignment = CTR
    ws.row_dimensions[row].height = 26
def section(ws, row, text, lastcol='J'):
    ws.merge_cells(f'A{row}:{lastcol}{row}')
    c = ws[f'A{row}']; c.value = text; c.font = SECT_FONT; c.fill = SECT_FILL; c.alignment = CTR
def note(ws, row, rows, text, lastcol='J'):
    ws.merge_cells(f'A{row}:{lastcol}{row+rows-1}')
    c = ws[f'A{row}']; c.value = text; c.font = NOTE_FONT; c.alignment = LEFTW
    ws.row_dimensions[row].height = 15*rows

# =====================================================================
# BUDGET IMPACT TAB
# =====================================================================
bi = wb['Budget Impact']

banner(bi, 60, 'JUN-DEC PHASED BUDGET PLAN (firm floor 110,000; top-ups staged separately, not in these tables)')

reframe = ("Reframe: drop 'no new budget' as the headline. The base rebuild and the reactivations are funded by redeploying current "
"waste, not by asking for more. On top of that, leadership has confirmed a firm Jun to Dec floor of 110,000 euros, which leadership "
"is adding to as the rebuilt engine proves it can absorb spend for return. Envelope facts: paid scope (search plus social) was about "
"242,600 euros for the year before this second-half plan; roughly 144,600 was committed Jan to May; the separate 14,800 trade-media "
"line is SEO and PR owned and out of paid scope. Full-year paid lands at about 254,600 (search about 199,000, social about 55,800). "
"The 45,000 non-recurrent LATAM line is closed and redistributed, so LATAM is a token line only. Definitions: Paid Search = Google "
"plus Microsoft/Bing; Paid Social = LinkedIn plus Meta.")
note(bi, 62, 4, reframe)

# ---- Market summary block (A-E) ----
section(bi, 67, 'MARKET SUMMARY, JUN-DEC (EUR)', lastcol='E')
sum_hdr = ['Market','Paid search','Paid social','Total Jun-Dec','Full-year (approx)']
for i,h in enumerate(sum_hdr):
    c = bi.cell(68, i+1, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = CTR; c.border = BORDER
# rows: market, search, social, fullyear-approx ; total = search+social formula
summary = [
    ('France', 32000, 11000, '~96,500'),
    ('UK',     27000,  8000, '~83,300'),
    ('Iberia', 20000,  9500, '~67,200'),
    ('LATAM (MX)', 1500, 1000, '~7,600'),
]
r = 69
for name, s, so, fy in summary:
    bi.cell(r,1,name).border = BORDER
    bi.cell(r,2,s).border = BORDER
    bi.cell(r,3,so).border = BORDER
    bi.cell(r,4).value = f'=B{r}+C{r}'; bi.cell(r,4).border = BORDER
    bi.cell(r,5,fy).border = BORDER
    r += 1
# All row (sums)
bi.cell(r,1,'All'); bi.cell(r,2).value=f'=SUM(B69:B{r-1})'; bi.cell(r,3).value=f'=SUM(C69:C{r-1})'
bi.cell(r,4).value=f'=SUM(D69:D{r-1})'; bi.cell(r,5).value=255  # placeholder replaced below
bi.cell(r,5).value=f'=96500+83300+67200+7600'
for c in range(1,6):
    cell=bi.cell(r,c); cell.font=ALL_FONT; cell.fill=ALL_FILL; cell.border=BORDER
SUMMARY_ALL_ROW = r  # 73
note(bi, r+1, 1, 'Paid Search = Google + Microsoft/Bing.  Paid Social = LinkedIn + Meta.', lastcol='E')

# ---- monthly block builder ----
months = ['Jun','Jul','Aug','Sep','Oct','Nov','Dec']
def month_block(start_row, title, data, fullyears, fy_formula):
    section(bi, start_row, title)
    hdr = ['Market'] + months + ['Jun-Dec','Full-year (approx)']
    for i,h in enumerate(hdr):
        c = bi.cell(start_row+1, i+1, h); c.font=HDR_FONT; c.fill=HDR_FILL; c.alignment=CTR; c.border=BORDER
    rr = start_row+2
    first = rr
    for (name, vals), fy in zip(data, fullyears):
        bi.cell(rr,1,name).border=BORDER
        for j,v in enumerate(vals):
            bi.cell(rr,2+j,v).border=BORDER
        bi.cell(rr,9).value=f'=SUM(B{rr}:H{rr})'; bi.cell(rr,9).border=BORDER  # I col = Jun-Dec
        bi.cell(rr,10,fy).border=BORDER  # J col full-year
        rr += 1
    last = rr-1
    # All row
    bi.cell(rr,1,'All').font=ALL_FONT
    for col in range(2,9):  # B..H months
        L=get_column_letter(col)
        bi.cell(rr,col).value=f'=SUM({L}{first}:{L}{last})'
    bi.cell(rr,9).value=f'=SUM(I{first}:I{last})'
    bi.cell(rr,10).value=fy_formula  # per-market full-year cells carry a "~" prefix, so sum the numeric components
    for col in range(1,11):
        cell=bi.cell(rr,col); cell.font=ALL_FONT; cell.fill=ALL_FILL; cell.border=BORDER
    note(bi, rr+1, 1, 'Paid Search = Google + Microsoft/Bing.  Paid Social = LinkedIn + Meta.')
    return rr  # All row index

search_data = [
    ('France',[5000,3000,3000,5000,6500,6500,3000]),
    ('UK',    [4000,2000,2000,4000,7000,7000,1000]),
    ('Iberia',[3000,1500,1500,3000,4000,4000,3000]),
    ('LATAM', [200,100,100,200,400,400,100]),
]
search_fy = ['~75,400','~65,100','~54,300','~4,000']
SEARCH_ALL_ROW = month_block(76, 'PAID SEARCH BY MONTH (EUR)', search_data, search_fy, '=75400+65100+54300+4000')

social_data = [
    ('France',[200,0,0,2200,4000,4000,600]),
    ('UK',    [150,0,0,1600,2900,2900,450]),
    ('Iberia',[200,0,0,2000,3500,3500,300]),
    ('LATAM', [0,0,0,200,400,400,0]),
]
social_fy = ['~21,100','~18,200','~12,900','~3,600']
SOCIAL_ALL_ROW = month_block(85, 'PAID SOCIAL BY MONTH (EUR)', social_data, social_fy, '=21100+18200+12900+3600')

# ---- two note rows ----
n1 = SOCIAL_ALL_ROW + 2
note1 = ("The 110,000 floor is firm. Search market totals are held; MX is trimmed to a token, and that budget plus the floor headroom "
"shifts into Iberia LinkedIn (social), so Iberia social now sits ahead of UK social for the period. Seasonal logic: summer floor "
"(near zero social in Jul and Aug), September ramp (a Sep lead closes in Oct or Nov), October and November peak, low December. Brand "
"defence (UK plus FR) is a named, protected slice within search (cheapest converter, was off five months). In-housing three external "
"retainers banks about 17,760 euros a year, kept as a line.")
note(bi, n1, 3, note1)
n2 = n1 + 3
note2 = ("Staged top-ups, held off the leadership slide and only in working docs and notes: an extra 30,000 to 40,000 likely soon, about "
"20,000 believed available for Q4 from unused budget, and 55,000 from a cancelled event earmarked for UK ideas. Deployed only on "
"confirmation, social growth-support and UK brand and comparison search first.")
note(bi, n2, 3, note2)

# widen month columns lightly so headers read
for col in ['B','C','D','E','F','G','H','I']:
    cur = bi.column_dimensions[col].width or 8
    if cur < 9:
        bi.column_dimensions[col].width = 9
bi.column_dimensions['J'].width = 16

print('Budget Impact done. Key rows -> summaryAll', SUMMARY_ALL_ROW, 'searchAll', SEARCH_ALL_ROW, 'socialAll', SOCIAL_ALL_ROW)

# =====================================================================
# REPORTING CADENCE TAB  (add KPI row)
# =====================================================================
rc = wb['Reporting Cadence']
rc.insert_rows(30, 1)  # blank separator pushed; expansion header moves 30->31
kpi_row = 29  # currently blank, sits right after re-baseline row 28
vals = ['Spend pace vs phased plan',
        'Actual month-to-date spend vs the Jun-Dec phased plan, by platform and market',
        'Platform spend + master Budget Impact tab', '-', 'YES', 'Weekly']
for i,v in enumerate(vals):
    c = rc.cell(kpi_row, i+1, v)
    c.alignment = Alignment(vertical='center', wrap_text=True, horizontal=('center' if i>=3 else 'left'))
print('Reporting Cadence KPI row added at', kpi_row)

# =====================================================================
# ACTION PLAN TAB  (answer-engine visibility + paid support menu pointer)
# =====================================================================
ap = wb['Action Plan']
r = ap.max_row + 1   # 46
ap.cell(r,1, 45)
ap.cell(r,2,'Q3')
ap.cell(r,3,'Answer-engine visibility: how Mapal shows up in AI answers and assistants. SEO, content and PR owned; paid only amplifies. Not a paid workstream.')
ap.cell(r,4,'SEO / Content / PR')
ap.cell(r,5,'MEDIUM')
ap.cell(r,6,'None')
ap.cell(r,7,'Q3 flag')
r += 1
ap.cell(r,1, 46)
ap.cell(r,2,'Ongoing')
ap.cell(r,3,'Paid support menu (see PPC and Nurture: The Plan): the standard set of proven plays for growth-team campaign requests, each with a setup and one judged metric. Use it to scope every growth-team paid request.')
ap.cell(r,4,'Paid + Growth')
ap.cell(r,5,'MEDIUM')
ap.cell(r,6,'None')
ap.cell(r,7,'Standard in use')
for rr in (r-1, r):
    for c in range(1,8):
        ap.cell(rr,c).alignment = Alignment(vertical='top', wrap_text=True)
print('Action Plan rows added at 46,47')

# =====================================================================
# COVER TAB  (v20 note)
# =====================================================================
cov = wb['Cover']
cov['C8'] = ("Last updated: 2026-06-03 (v20). Added the Jun to Dec phased budget plan to the Budget Impact tab: a firm 110,000 euro "
"floor, market summary plus paid search and paid social by month (France, UK, Iberia, LATAM), with the All rows as live SUM formulas "
"(Jun-Dec 80,500 search and 29,500 social; full-year about 254,600). Reframed away from 'no new budget' to a disciplined base funded "
"by cutting waste, with staged top-ups held off-plan. MX trimmed into Iberia LinkedIn so Iberia social now sits ahead of UK; the "
"45,000 non-recurrent LATAM line is closed and redistributed. Added a Spend-pace-vs-phased-plan KPI (Reporting Cadence) and Q3 "
"answer-engine visibility plus a paid support menu pointer (Action Plan). Locked Numbers tab unchanged.")
cov['C8'].alignment = Alignment(wrap_text=True, vertical='top')
print('Cover updated')

# =====================================================================
# CHANGE LOG TAB  (v20 entry)
# =====================================================================
cl = wb['Change Log']
r = cl.max_row + 1  # 25
cl.cell(r,1,'23')
cl.cell(r,2,'2026-06-03')
cl.cell(r,3,'Adrian')
cl.cell(r,4, ("Master sheet v20. Added the JUN-DEC PHASED BUDGET PLAN to the Budget Impact tab below the existing content: a firm 110,000 "
"euro floor, a market summary block, and paid-search and paid-social by-month blocks (Jun to Dec) for France, UK, Iberia and LATAM, "
"with Jun-Dec columns and All rows as live SUM formulas and hardcoded full-year approximations. All rows compute to 80,500 search and "
"29,500 social for Jun-Dec (110,000 total) and about 198,800 search, 55,800 social, 254,600 total full-year. Reframed away from 'no "
"new budget' to a disciplined base funded by cutting waste that leadership is adding to; staged top-ups held off-plan. MX trimmed to a "
"token and shifted into Iberia LinkedIn so Iberia social now sits ahead of UK social; the 45,000 non-recurrent LATAM line is closed "
"and redistributed. Added a Spend-pace-vs-phased-plan KPI to Reporting Cadence and two Action Plan rows (Q3 answer-engine visibility, "
"SEO/content/PR owned with paid amplifying; and a pointer to the paid support menu in The Plan as the standard for growth-team "
"requests). Locked Numbers tab untouched."))
cl.cell(r,4).alignment = Alignment(wrap_text=True, vertical='top')
print('Change Log updated at row', r)

# force recalc on load so LibreOffice/Excel evaluate the SUM formulas
wb.calculation.fullCalcOnLoad = True
import os
os.makedirs('FINAL_v_current', exist_ok=True)
wb.save(OUT)
print('SAVED', OUT)
